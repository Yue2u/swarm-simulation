#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Собирает отчёт из набора markdown-файлов в один .docx с картинками.

Скрипт читает все файлы docs/report/NN-*.md в порядке номеров, разбирает
поддерживаемое подмножество markdown (заголовки, абзацы, списки, таблицы,
картинки, блоки кода) и складывает всё в один документ Word через python-docx.

Запуск:  python3 docs/report/build_report.py
Результат:  docs/report/boids-report.docx
"""

from __future__ import annotations

import glob
import os
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

REPORT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT = os.path.join(REPORT_DIR, "boids-report.docx")

BODY_FONT = "Calibri"
CODE_FONT = "Consolas"
CODE_FILL = "F2F2F2"
GREY = RGBColor(0x66, 0x66, 0x66)
CODE_COLOR = RGBColor(0x8B, 0x1A, 0x1A)

# ---------------------------------------------------------------------------------------------
# Разбор markdown
# ---------------------------------------------------------------------------------------------

TOKEN_RE = re.compile(
    r"(\*\*.+?\*\*|`[^`]+`|\*[^*]+?\*|_[^_]+_|\[[^\]]+?\]\([^)]+?\))"
)


def parse_blocks(text: str):
    """Возвращает список блоков (kind, data)."""
    lines = text.split("\n")
    blocks = []
    i = 0
    n = len(lines)
    while i < n:
        raw = lines[i]
        s = raw.strip()

        if s.startswith("```"):
            i += 1
            code = []
            while i < n and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            i += 1  # закрывающий забор
            blocks.append(("code", "\n".join(code)))
            continue

        if s == "---":
            blocks.append(("hr", None))
            i += 1
            continue

        if s.startswith("|") and s.count("|") >= 2:
            rows = []
            while i < n and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                rows.append(cells)
                i += 1
            # выкидываем строку-разделитель |---|---|
            rows = [
                r
                for r in rows
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in r if c != "")
            ]
            blocks.append(("table", rows))
            continue

        if s.startswith("!["):
            m = re.match(r"!\[(.*?)\]\((.*?)\)", s)
            if m:
                blocks.append(("image", (m.group(1), m.group(2))))
                i += 1
                continue

        m = re.match(r"^(#{1,6})\s+(.*)$", s)
        if m:
            blocks.append(("h", (len(m.group(1)), m.group(2).strip())))
            i += 1
            continue

        if re.match(r"^[-*]\s+", s):
            blocks.append(("bullet", re.sub(r"^[-*]\s+", "", s)))
            i += 1
            continue

        m = re.match(r"^(\d+)\.\s+(.*)$", s)
        if m:
            blocks.append(("ordered", m.group(2)))
            i += 1
            continue

        if s.startswith(">"):
            blocks.append(("quote", s.lstrip("> ").strip()))
            i += 1
            continue

        if s == "":
            i += 1
            continue

        para = [s]
        i += 1
        while i < n:
            nxt = lines[i].strip()
            if (
                nxt == ""
                or nxt.startswith(("#", "|", "!", ">", "```", "- ", "* "))
                or re.match(r"^\d+\.\s", nxt)
                or nxt == "---"
            ):
                break
            para.append(nxt)
            i += 1
        blocks.append(("p", " ".join(para)))
    return blocks


# ---------------------------------------------------------------------------------------------
# Оформление
# ---------------------------------------------------------------------------------------------


def add_runs(paragraph, text: str):
    """Добавляет текст с инлайн-разметкой: **жирный**, *курсив*, `код`, [ссылка](url)."""
    for part in TOKEN_RE.split(text):
        if part == "":
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            r = paragraph.add_run(part[2:-2])
            r.bold = True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            r = paragraph.add_run(part[1:-1])
            r.font.name = CODE_FONT
            r.font.color.rgb = CODE_COLOR
        elif part.startswith("[") and "](" in part:
            label = part[1 : part.index("](")]
            r = paragraph.add_run(label)
            r.italic = True
        elif (part.startswith("*") and part.endswith("*") and len(part) > 2) or (
            part.startswith("_") and part.endswith("_") and len(part) > 2
        ):
            r = paragraph.add_run(part[1:-1])
            r.italic = True
        else:
            paragraph.add_run(part)
    return paragraph


def shade(paragraph, fill: str):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)


def add_page_number(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def add_toc(doc):
    p = doc.add_paragraph()
    run = p.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "Оглавление: обновите поле (F9) в Word."
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for el in (begin, instr, sep, placeholder, end):
        run._r.append(el)


def add_image(doc, alt, path):
    full = path if os.path.isabs(path) else os.path.join(REPORT_DIR, path)
    if not os.path.exists(full):
        print(f"  ! нет картинки: {full}")
        return
    doc.add_picture(full, width=Cm(15.0))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if alt:
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cap.add_run(alt)
        r.italic = True
        r.font.size = Pt(9)
        r.font.color.rgb = GREY


def add_table(doc, rows):
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=0, cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    try:
        table.style = "Light Grid Accent 1"
    except KeyError:
        table.style = "Table Grid"
    for ri, row in enumerate(rows):
        cells = table.add_row().cells
        for ci in range(cols):
            text = row[ci] if ci < len(row) else ""
            cell = cells[ci]
            cell.text = ""
            p = cell.paragraphs[0]
            add_runs(p, text)
            for r in p.runs:
                r.font.size = Pt(9.5)
                if ri == 0:
                    r.bold = True
    doc.add_paragraph()


# ---------------------------------------------------------------------------------------------
# Сборка
# ---------------------------------------------------------------------------------------------


def main():
    files = sorted(glob.glob(os.path.join(REPORT_DIR, "[0-9]*-*.md")))
    if not files:
        raise SystemExit("Не найдено ни одного NN-*.md в " + REPORT_DIR)
    print(f"Файлов: {len(files)}")

    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_number(footer)

    first_h1_done = False
    in_title = False

    for path in files:
        name = os.path.basename(path)
        print(f"  + {name}")
        with open(path, encoding="utf-8") as fh:
            text = fh.read()
        blocks = parse_blocks(text)
        in_title = name.startswith("00")

        for kind, data in blocks:
            if kind == "h":
                level, title = data
                if level == 1 and not first_h1_done:
                    first_h1_done = True
                    p = doc.add_paragraph(title, style="Title")
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    continue
                if level == 1:
                    doc.add_page_break()
                    doc.add_heading(title, level=1)
                else:
                    doc.add_heading(title, level=min(level, 3))
                continue

            if kind == "p":
                p = doc.add_paragraph()
                add_runs(p, data)
                if in_title:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                continue

            if kind == "bullet":
                p = doc.add_paragraph(style="List Bullet")
                add_runs(p, data)
                continue

            if kind == "ordered":
                p = doc.add_paragraph(style="List Number")
                add_runs(p, data)
                continue

            if kind == "quote":
                p = doc.add_paragraph(style="Intense Quote")
                add_runs(p, data)
                continue

            if kind == "table":
                add_table(doc, data)
                continue

            if kind == "image":
                add_image(doc, data[0], data[1])
                continue

            if kind == "code":
                for line in str(data).split("\n"):
                    p = doc.add_paragraph()
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.line_spacing = 1.0
                    r = p.add_run(line if line else " ")
                    r.font.name = CODE_FONT
                    r.font.size = Pt(9)
                    shade(p, CODE_FILL)
                doc.add_paragraph()
                continue

            if kind == "hr":
                if in_title and not first_h1_done:
                    continue
                if in_title:
                    doc.add_page_break()
                    in_title = False
                    add_toc(doc)
                    doc.add_page_break()
                    continue
                p = doc.add_paragraph()
                pPr = p._p.get_or_add_pPr()
                pbdr = OxmlElement("w:pBdr")
                bottom = OxmlElement("w:bottom")
                bottom.set(qn("w:val"), "single")
                bottom.set(qn("w:sz"), "6")
                bottom.set(qn("w:space"), "1")
                bottom.set(qn("w:color"), "BBBBBB")
                pbdr.append(bottom)
                pPr.append(pbdr)
                continue

    doc.save(OUTPUT)
    print(f"Готово: {OUTPUT}")


if __name__ == "__main__":
    main()
